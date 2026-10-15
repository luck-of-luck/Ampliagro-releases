import json
import logging
import math
import random
import traceback
from collections import defaultdict
from datetime import date, datetime, timedelta

import requests
from django.conf import settings
from django.contrib.auth import authenticate, login
from django.contrib.auth.models import User
from django.core.cache import cache
from django.db import transaction
from django.db.models import Avg, Count, F, Q, Sum
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.utils.dateparse import parse_date
from drf_spectacular.utils import OpenApiParameter, extend_schema
from rest_framework import permissions, status, viewsets
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from app.api.serializers import *
from app.hierarquia import *
from app.models import *

from .utils import *

#: Os geradores de ficheiro, por formato de `RelatorioTemplate.formato`.
#:
#: VAZIO de propósito, e é o ponto. `configuracao` é um JSON livre sem esquema:
#: nada define o que é uma configuração válida, e um gerador escrito agora seria
#: um gerador para um formato que ninguém acordou. Enquanto esta tabela estiver
#: vazia, `gerar_relatorio` responde 501 com o motivo — que é informação — em vez
#: de devolver `processando` para sempre, que é uma mentira.
#:
#: Para exportar já hoje, e com o tenancy aplicado, existem
#: `/api/dashboard-export/csv/` e `/api/dashboard-export/excel/`.
EXPORTADORES: dict[str, object] = {}


class RelatorioGeradoViewSet(viewsets.ModelViewSet):
    filterset_fields = ["template"]
    ordering_fields = ["criado_em"]
    serializer_class = RelatorioGeradoSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return RelatorioGerado.objects.filter(
            template__user=self.request.user
        ).select_related("template")


class RelatorioTemplateViewSet(viewsets.ModelViewSet):
    search_fields = ["nome"]
    serializer_class = RelatorioTemplateSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        return RelatorioTemplate.objects.filter(user=self.request.user)


@extend_schema(responses=SPASimpleSerializer)
@api_view(["POST"])
@permission_classes([IsAuthenticated])
def gerar_relatorio(request):
    """Generate report from template.

    O QUE MUDOU, E POR QUÊ
    ======================
    A versão anterior criava a linha com `status="processando"` e devolvia. Não
    havia task Celery, não havia geração síncrona — apesar do comentário dizer
    "For now: generate synchronously". O ficheiro nunca era escrito, e
    `relatorio_status` respondia `processando` e `arquivo: null` **para sempre**.

    É a falha que o próprio repo nomeia: "um botão que parece funcionar e não faz
    nada". E aqui era pior que um botão: era um botão que prometia um ficheiro e
    dizia ao produtor que estava a trabalhar.

    A correcção NÃO é inventar um gerador. `RelatorioTemplate.configuracao` é um
    JSON livre, sem esquema — não há contrato que diga o que uma configuração
    válida é, e inventar um agora seria escrever um gerador para um formato que
    ninguém definiu. O que se faz é o honesto: **ou gera, ou diz que não
    gera.** Um `status="erro"` com o motivo é informação; um `processando`
    eterno é uma mentira que o produtor não tem como desmaiar.

    `EXPORTADORES` é o sítio onde um gerador se liga quando existir. Está vazio
    de propósito, e o teste `test_gerar_relatorio_nunca_deixa_processando` fecha
    a porta ao modo antigo.
    """
    from app.models import RelatorioGerado, RelatorioTemplate

    template_id = request.data.get("template_id")
    params = request.data.get("params", {})
    template = RelatorioTemplate.objects.filter(
        pk=template_id, user=request.user
    ).first()
    if not template:
        return Response({"detail": "Template não encontrado"}, status=404)

    rel = RelatorioGerado.objects.create(
        template=template, user=request.user, parametros=params, status="processando"
    )

    gerador = EXPORTADORES.get(template.formato)
    if gerador is None:
        rel.status = "erro"
        rel.erro_msg = (
            f"Não há gerador para o formato '{template.formato}'. "
            "O relatório NÃO foi gerado — o que existe em 'RelatorioGerado' é o "
            "pedido, não o ficheiro. Para exportar já hoje, use "
            "/api/dashboard-export/csv/ ou /api/dashboard-export/excel/."
        )
        rel.save(update_fields=["status", "erro_msg", "concluido_em"])
        return Response(
            {
                "relatorio_id": rel.pk,
                "status": "erro",
                "erro": rel.erro_msg,
            },
            status=501,
        )

    try:
        rel.arquivo = gerador(template, params)
        rel.status = "concluido"
        rel.save(update_fields=["status", "arquivo", "concluido_em"])
    except Exception as exc:  # noqa: BLE001 — o erro VAI para a linha, não para o log
        rel.status = "erro"
        rel.erro_msg = str(exc)[:2000]
        rel.save(update_fields=["status", "erro_msg", "concluido_em"])

    return Response({"relatorio_id": rel.pk, "status": rel.status})


@extend_schema(responses=SPASimpleSerializer)
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def relatorio_status(request, pk):
    """Check report generation status"""
    from app.models import RelatorioGerado

    rel = RelatorioGerado.objects.filter(pk=pk, user=request.user).first()
    if not rel:
        return Response({"detail": "Não encontrado"}, status=404)
    return Response(
        {
            "status": rel.status,
            "arquivo": rel.arquivo.url if rel.arquivo else None,
            "erro": rel.erro_msg,
        }
    )


@extend_schema(responses=SPASimpleSerializer)
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def export_dashboard_csv(request):
    """Export dashboard data as CSV"""
    import csv

    from django.http import HttpResponse

    props = get_propriedade_optimizada(request.user)
    if not props:
        return HttpResponse("Nenhuma propriedade", status=400)
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = 'attachment; filename="dashboard_export.csv"'
    response.write("\ufeff")  # BOM for Excel
    writer = csv.writer(response, delimiter=";")
    writer.writerow(["Tipo", "Campo", "Valor"])
    # Safras
    for s in (
        Safra.objects.filter(propriedade=props)
        .select_related("cultura")
        .order_by("-ano")[:20]
    ):
        writer.writerow(["Safra", "Cultura", s.cultura.nome if s.cultura else ""])
        writer.writerow(["", "Ano", s.ano])
        writer.writerow(["", "Área (ha)", s.area_plantada])
        writer.writerow(["", "Produção (sc)", s.producao_total])
        writer.writerow(["", "Produtividade (sc/ha)", s.produtividade])
    # Financeiro
    for l in LancamentoFinanceiro.objects.filter(propriedade=props).order_by("-data")[
        :50
    ]:
        writer.writerow(["Financeiro", l.categoria, f"{l.tipo}: R$ {l.valor}"])
    # Estoque
    for e in EstoqueItem.objects.filter(propriedade=props):
        writer.writerow(["Estoque", e.nome, f"{e.quantidade} {e.unidade}"])
    return response


@extend_schema(responses=SPASimpleSerializer)
@api_view(["GET"])
@permission_classes([IsAuthenticated])
def export_dashboard_excel(request):
    """Export dashboard data as Excel"""
    import openpyxl
    from django.http import HttpResponse

    props = get_propriedade_optimizada(request.user)
    if not props:
        return HttpResponse("Nenhuma propriedade", status=400)
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Safras"
    ws1.append(
        [
            "Cultura",
            "Ano",
            "Área (ha)",
            "Produção (sc)",
            "Produtividade (sc/ha)",
            "Receita (R$)",
        ]
    )
    for s in (
        Safra.objects.filter(propriedade=props)
        .select_related("cultura")
        .order_by("-ano")
    ):
        ws1.append(
            [
                s.cultura.nome if s.cultura else "",
                s.ano,
                float(s.area_plantada or 0),
                float(s.producao_total or 0),
                float(s.produtividade or 0),
                float(s.receita_estimada or 0),
            ]
        )
    ws2 = wb.create_sheet("Financeiro")
    ws2.append(["Data", "Tipo", "Categoria", "Descrição", "Valor (R$)", "Pago"])
    for l in LancamentoFinanceiro.objects.filter(propriedade=props).order_by("-data")[
        :100
    ]:
        ws2.append(
            [
                l.data,
                l.tipo,
                l.categoria,
                l.descricao,
                float(l.valor),
                "Sim" if l.pago else "Não",
            ]
        )
    ws3 = wb.create_sheet("Estoque")
    ws3.append(["Item", "Categoria", "Quantidade", "Unidade", "Local", "Validade"])
    for e in EstoqueItem.objects.filter(propriedade=props):
        ws3.append(
            [
                e.nome,
                e.categoria,
                float(e.quantidade or 0),
                e.unidade,
                e.local,
                e.data_validade or "",
            ]
        )
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="dashboard_export.xlsx"'
    wb.save(response)
    return response
